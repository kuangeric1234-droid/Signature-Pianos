"""Build the Word letterhead (A4) from the same marks and measurements as letterhead.html.

Run from the repo root:  python brand/src/build_docx.py
Install Jost (brand/Fonts/Jost/Static) before typing letters, or Word substitutes its default.
"""
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor
from PIL import Image

ROOT = Path(__file__).resolve().parent.parent
LH = ROOT / 'Communication assets' / 'Letterhead'
LOGO = LH / 'src' / 'letterhead-logo.png'
MONO = ROOT / 'Brand Assets' / '03. Monogram' / '01. Ink' / 'SignaturePianos_Monogram_Ink.png'
FELT = LH / 'src' / 'felt-rule.png'

INK = RGBColor(0x16, 0x22, 0x2E)
INK_SOFT = RGBColor(0x3A, 0x48, 0x55)


def spacing(run, twentieths):
    """Letter-spacing in twentieths of a point (Word's w:spacing)."""
    rpr = run._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing')
    sp.set(qn('w:val'), str(twentieths))
    rpr.append(sp)


def font(run, size, colour, name='Jost', bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), name)
    run.font.size = Pt(size)
    run.font.color.rgb = colour
    run.font.bold = bold


def main():
    # the key-slip felt rule as an image, 10 x 0.5 mm
    Image.new('RGB', (400, 20), (0x4A, 0x15, 0x20)).save(FELT)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(28)
    sec.top_margin, sec.bottom_margin = Mm(62), Mm(52)
    sec.header_distance, sec.footer_distance = Mm(17), Mm(12)

    normal = doc.styles['Normal']
    normal.font.name = 'Jost'
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), 'Jost')
    normal.font.size = Pt(9.8)
    normal.font.color.rgb = INK_SOFT
    pf = normal.paragraph_format
    pf.space_after = Pt(7.5)
    pf.line_spacing = 1.3

    # header: the lockup, then the felt rule
    hp = sec.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hp.add_run().add_picture(str(LOGO), width=Mm(50))
    hp.paragraph_format.space_after = Pt(4)
    hf = sec.header.add_paragraph()
    hf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hf.add_run().add_picture(str(FELT), width=Mm(10), height=Mm(0.5))

    # footer: the monogram, the address in tracked capitals
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fp.paragraph_format.space_after = Pt(9)
    fp.add_run().add_picture(str(MONO), width=Mm(5))
    for text in ('63 Blackburn Road, Mount Waverley VIC 3149',
                 '0479 128 955  ·  info@signaturepianos.com.au  ·  signaturepianos.com.au'):
        p = sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text.upper())
        font(r, 6.4, INK_SOFT)
        spacing(r, 26)

    # a starter letter in the house voice; replace the text, keep the styles
    lines = [
        ('19 September 2026', False), ('', False),
        ('Ms Sarah Nguyen\n12 Station Street\nBox Hill VIC 3128', False), ('', False),
        ('Dear Sarah', False),
        ('Your Yamaha U3H is ready for delivery', True),
        ('Thank you for choosing your piano with us. The U3H you played in the showroom was hand-picked by our buyer in Japan, brought straight to Melbourne and checked in our Mount Waverley workshop before it went on the floor.', False),
        ('Delivery is white-glove and tracked from start to finish. You can pick the day and time window that suits you in your customer portal; we place the piano where you choose and send a photo the moment it is in position.', False),
        ('Your 10-year warranty certificate is emailed as soon as delivery is confirmed. Your first tuning is included and is booked for three to four weeks after delivery, once the piano has settled into your room.', False),
        ('If anything about the piano is not right, call me directly on 0479 128 955.', False),
        ('Kind regards', False), ('', False), ('', False),
        ('Eric\nSignature Pianos', False),
    ]
    for text, strong in lines:
        p = doc.add_paragraph()
        r = p.add_run(text)
        if strong:
            font(r, 9.8, INK, bold=False)
            r.font.name = 'Jost Medium'
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:ascii'), 'Jost Medium')
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:hAnsi'), 'Jost Medium')

    doc.core_properties.title = 'Signature Pianos Letterhead'
    doc.core_properties.author = 'Signature Pianos'
    out = LH / 'SignaturePianos_Letterhead.docx'
    doc.save(out)
    print('wrote', out.relative_to(ROOT))


if __name__ == '__main__':
    main()
